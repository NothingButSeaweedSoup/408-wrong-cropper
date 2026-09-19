"""7.4 word_builder：把排版方案写成 docx。

要点（AGENTS.md 7.4 / 11）：
- A4、页边距 2cm；
- 图片宽度统一 image_width_cm，高度等比缩放；
- 跨页题的多张图片连续插入，中间不加分页符；
- 用 keepNext(w:keepNext) + keepLines(w:keepLines) 保证一道题的图片不被拆开，
  换页通过"段前分页"(w:pageBreakBefore) 实现，而不是插入分页符空段。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .. import config
from .layout import Entry, Piece

CJK_FONT = "宋体"


def export_filename(when: datetime | None = None) -> str:
    when = when or datetime.now()
    return f"408错题本_{when:%Y-%m-%d_%H%M}.docx"


def _setup_section(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(config.PAGE_W_CM)
    section.page_height = Cm(config.PAGE_H_CM)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(config.MARGIN_CM))


def _setup_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), CJK_FONT)


def _image_paragraph(doc: Document, path: Path, width_cm: float, *, keep_next: bool):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(config.SPACING_PT_PER_GAP)
    fmt.keep_together = True  # w:keepLines
    fmt.keep_with_next = keep_next  # w:keepNext
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return p


def _caption_paragraph(doc: Document, text: str):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(2)
    fmt.keep_together = True
    fmt.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    return p


def _note_paragraphs(doc: Document, count: int) -> None:
    for _ in range(max(0, count)):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        run = p.add_run(" ")
        run.font.size = Pt(config.NOTE_LINE_CM / (2.54 / 72.0))


def materialize_pieces(
    entries: list[Entry],
    page_paths: dict[int, dict[int, Path]],
    out_root: Path,
) -> None:
    """把"按小问拆出来的段"实际裁成图片文件（整块的题目本来就有图，跳过）。

    这样做而不是在生成 Word 时现裁，是为了让裁剪产物可复用、可缓存。
    """
    from .cropper import crop_block

    for entry in entries:
        question = entry.question
        bbox = question.get("bbox") or {}
        x0 = int(bbox.get("x0") or 0)
        x1 = bbox.get("x1")
        paper_id = int(question.get("paper_id") or 0)
        year = int(question.get("year") or 0)
        qno = int(question.get("question_no") or 0)
        pages = page_paths.get(paper_id, {})
        for index, piece in enumerate(entry.pieces, start=1):
            if piece.path:
                continue
            page_path = pages.get(piece.page_no)
            if page_path is None:
                continue
            out = out_root / str(paper_id) / f"{year}_q{qno}_p{piece.page_no}_{index}.png"
            crop_block(page_path, piece.y0, piece.y1, out, x0=x0, x1=int(x1) if x1 else None)
            piece.path = config.rel_path(out)


def build(
    entries: list[Entry],
    out_path: str | Path,
    *,
    image_width_cm: float | None = None,
    with_caption: bool = False,
    note_lines: int = 0,
    title: str | None = None,
) -> dict:
    """生成 docx，返回统计信息。"""
    image_width_cm = float(image_width_cm or config.IMAGE_WIDTH_CM)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _setup_section(doc)
    _setup_styles(doc)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = p.paragraph_format
        fmt.space_after = Pt(10)
        fmt.keep_with_next = True
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    root = config.ROOT_DIR
    image_count = 0
    missing: list[str] = []

    for entry in entries:
        question = entry.question
        pieces: list[Piece] = entry.pieces
        if not pieces:
            continue

        first_marker = True

        if with_caption:
            year = question.get("year")
            qno = question.get("question_no")
            text = f"{year} 年第 {qno} 题" if year else f"第 {qno} 题"
            if entry.sub_label:
                text += f" {entry.sub_label}"
            p = _caption_paragraph(doc, text)
            if entry.new_page:
                p.paragraph_format.page_break_before = True
            first_marker = False

        for i, piece in enumerate(pieces):
            if not piece.path:
                missing.append(f"q{question.get('question_no')} p{piece.page_no}")
                continue
            path = Path(piece.path)
            if not path.is_absolute():
                path = root / path
            if not path.exists():
                missing.append(str(path))
                continue
            keep_next = i < len(pieces) - 1
            p = _image_paragraph(doc, path, image_width_cm, keep_next=keep_next)
            if first_marker and entry.new_page:
                p.paragraph_format.page_break_before = True
            first_marker = False
            image_count += 1

        _note_paragraphs(doc, note_lines)

    doc.save(out_path)
    return {
        "image_count": image_count,
        "entry_count": len(entries),
        "missing": missing,
    }
